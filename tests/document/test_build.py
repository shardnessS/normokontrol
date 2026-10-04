"""Invariants of the generated docx (до этапа 7 — проверки через python-docx)."""

import shutil
import zipfile
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from normokontrol.bibliography.names import NBSP
from normokontrol.cli import EXIT_ERROR, EXIT_OK, run
from normokontrol.cli import document as document_cli
from normokontrol.document.build import build_file
from normokontrol.document.markdown_in import DocumentError
from normokontrol.document.math_omml import M_NS

ROOT = Path(__file__).resolve().parents[2]
EXAMPLE = ROOT / "examples" / "coursework"


@pytest.fixture
def example(tmp_path: Path) -> Path:
    """Copy of the example (Cyrillic file name) in a temp folder."""
    target = tmp_path / "курсовая"
    shutil.copytree(EXAMPLE, target, ignore=shutil.ignore_patterns("*.docx"))
    return target / "черновик.md"


@pytest.fixture
def built(example: Path) -> docx.document.Document:
    result = build_file(example)
    assert result.output == example.with_name("черновик_gost.docx")
    return docx.Document(str(result.output))


def body_paragraphs(document: docx.document.Document) -> list[str]:
    return [p.text.replace(NBSP, " ") for p in document.paragraphs]


def test_page_setup(built: docx.document.Document) -> None:
    section = built.sections[0]
    assert round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297
    margins = [section.left_margin, section.right_margin, section.top_margin, section.bottom_margin]
    assert [round(m.mm) for m in margins] == [30, 15, 20, 20]


def test_normal_style(built: docx.document.Document) -> None:
    normal = built.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 14
    fmt = normal.paragraph_format
    assert fmt.line_spacing == 1.5
    assert round(fmt.first_line_indent.cm, 2) == 1.25
    assert fmt.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert fmt.space_before.pt == 0 and fmt.space_after.pt == 0
    rfonts = normal.element.rPr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:asciiTheme")) is None  # тема шрифтов не перекрывает Times New Roman


def test_headings(built: docx.document.Document) -> None:
    by_text = {p.text: p for p in built.paragraphs}
    intro = by_text["ВВЕДЕНИЕ"]
    assert intro.style.name == "Heading 1"
    assert intro.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert intro.paragraph_format.page_break_before is True
    section = by_text["1 Анализ предметной области"]
    assert section.paragraph_format.page_break_before is True
    assert round(section.paragraph_format.first_line_indent.cm, 2) == 1.25
    subsection = by_text["1.1 Обзор существующих решений"]
    assert subsection.style.name == "Heading 2"
    assert subsection.paragraph_format.page_break_before is False
    heading_font = built.styles["Heading 1"].font
    assert heading_font.bold is True and heading_font.color.rgb == docx.shared.RGBColor(0, 0, 0)
    assert not any(
        text.endswith(".") for text in by_text if text and by_text[text].style.name.startswith("Heading")
    )


def test_captions_and_references(built: docx.document.Document) -> None:
    texts = body_paragraphs(built)
    assert "Рисунок 1 — Архитектура системы учёта заявок" in texts
    assert "Таблица 1 — Сравнение систем учёта заявок" in texts
    assert any("Как показано на рисунке 1, разрабатываемая" in text for text in texts)
    figure_caption = next(p for p in built.paragraphs if p.text.startswith("Рисунок 1"))
    assert figure_caption.alignment == WD_ALIGN_PARAGRAPH.CENTER
    table_caption = next(p for p in built.paragraphs if p.text.startswith("Таблица 1"))
    assert table_caption.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert not any("[@" in text for text in texts)


def test_table(built: docx.document.Document) -> None:
    table = built.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Решение", "Назначение", "Стоимость, руб. в месяц"]
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.rows[1].cells[0].paragraphs[0].runs[0].font.size is None  # размер — из стиля
    assert built.styles["GOST Table Text"].font.size.pt == 12


def test_formulas(built: docx.document.Document) -> None:
    numbered = [p for p in built.paragraphs if p._p.find(f"{{{M_NS}}}oMath") is not None]
    assert len(numbered) == 2
    assert [p.text.strip() for p in numbered] == ["(1)", "(2)"]
    tab_stops = numbered[0].paragraph_format.tab_stops
    assert len(tab_stops) == 2
    explanation = next(p for p in built.paragraphs if p.text.startswith("где T"))
    assert explanation._p.find(f".//{qn('w:br')}") is not None  # каждое обозначение с новой строки


def test_bibliography_and_appendix(built: docx.document.Document) -> None:
    texts = body_paragraphs(built)
    start = texts.index("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    assert texts[start + 1].startswith("1. Варламова, Л. Н. Управление документацией")
    assert texts[start + 6].startswith("6. Порядок присвоения номера ISBN")
    appendix = next(p for p in built.paragraphs if p.text.startswith("ПРИЛОЖЕНИЕ А"))
    assert appendix.text == "ПРИЛОЖЕНИЕ А\nЛистинг модуля регистрации заявок"
    assert appendix.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert appendix.paragraph_format.page_break_before is True
    assert any("приведён в приложении А" in text for text in texts)


def test_fields(built: docx.document.Document) -> None:
    body = built.element.body
    instructions = [node.text for node in body.iter(qn("w:instrText"))]
    assert any("TOC" in text for text in instructions)
    footer = built.sections[0].footer.paragraphs[0]
    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert any("PAGE" in node.text for node in footer._p.iter(qn("w:instrText")))
    assert built.settings.element.find(qn("w:updateFields")) is not None


def test_build_is_deterministic(example: Path, tmp_path: Path) -> None:
    first = build_file(example, output=tmp_path / "a.docx").output.read_bytes()
    second = build_file(example, output=tmp_path / "b.docx").output.read_bytes()
    assert first == second
    with zipfile.ZipFile(tmp_path / "a.docx") as zf:
        assert {info.date_time for info in zf.infolist()} == {(2000, 1, 1, 0, 0, 0)}


def test_source_is_never_overwritten(example: Path) -> None:
    original = example.read_bytes()
    with pytest.raises(DocumentError, match="не может перезаписать исходный файл"):
        build_file(example, output=example)
    assert example.read_bytes() == original


def test_missing_image_and_bad_formula(tmp_path: Path) -> None:
    draft = tmp_path / "d.md"
    draft.write_text("# Раздел\n\n![Нет](missing.png)\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="Не найден файл рисунка") as info:
        build_file(draft)
    assert info.value.location == "d.md:3"
    draft.write_text("# Раздел\n\n$$ \\foo $$\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="«\\\\foo» не поддерживается"):
        build_file(draft)
    (tmp_path / "pic.svg").write_text("<svg/>", encoding="utf-8")
    draft.write_text("# Раздел\n\n![Векторный](pic.svg)\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="Формат рисунка"):
        build_file(draft)


def test_sources_inline_and_errors(tmp_path: Path) -> None:
    draft = tmp_path / "d.md"
    draft.write_text(
        "---\nsources:\n  - {type: web, title: Сайт, url: 'https://x.ru', accessed: 2026-01-01}\n---\n"
        "# ВВЕДЕНИЕ\nТекст.\n",
        encoding="utf-8",
    )
    with pytest.raises(DocumentError, match="нет id"):
        build_file(draft)
    draft.write_text("---\nsources: [просто текст]\n---\n# ВВЕДЕНИЕ\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="Не удалось получить данные источников"):
        build_file(draft)
    (tmp_path / "s.yaml").write_text("source: 1\n", encoding="utf-8")
    draft.write_text("---\nsources: s.yaml\n---\n# ВВЕДЕНИЕ\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="ожидался список источников"):
        build_file(draft)


def test_cli_summary(example: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert run(document_cli.main, [str(example)]) == EXIT_OK
    out = capsys.readouterr().out
    assert "черновик_gost.docx" in out
    assert "разделов: 2, рисунков: 1, таблиц: 1, формул: 2, приложений: 1, источников: 6" in out
    assert run(document_cli.main, [str(example.with_name("нет.md"))]) == EXIT_ERROR
    assert "document_invalid" in capsys.readouterr().err
