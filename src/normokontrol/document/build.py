"""Prepared blocks → .docx with python-docx. Output is byte-for-byte reproducible."""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
from docx.document import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.image.exceptions import UnrecognizedImageError
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from normokontrol.bibliography.inputs import read_data, resolve
from normokontrol.bibliography.models import Source
from normokontrol.document import ooxml, styles
from normokontrol.document.markdown_in import DocumentError, ParsedDocument, parse_file
from normokontrol.document.math_omml import MathError, latex_to_omml
from normokontrol.document.preprocess import (
    OutBibliography,
    OutBlock,
    OutCode,
    OutFigure,
    OutFormula,
    OutHeading,
    OutList,
    OutParagraph,
    OutTable,
    Prepared,
    prepare,
)
from normokontrol.presets import loader
from normokontrol.presets.schema import Preset

# Фиксированные дата и время: одинаковый вход даёт побайтово одинаковый docx.
FIXED_TIME = dt.datetime(2000, 1, 1, tzinfo=dt.UTC)
ZIP_DATE = (2000, 1, 1, 0, 0, 0)
TOC_FIELD = 'TOC \\o "1-3" \\h \\z \\u'
TOC_PLACEHOLDER = "Оглавление обновится при открытии в Word (если нет — правый клик → «Обновить поле»)."
_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


@dataclass
class BuildResult:
    output: Path
    stats: dict[str, int]
    warnings: list[str] = field(default_factory=list)


def build_file(input_path: Path, *, preset_id: str | None = None, output: Path | None = None) -> BuildResult:
    """Markdown file → docx next to it (`<name>_gost.docx`). The source file is never overwritten."""
    document = parse_file(input_path)
    preset = loader.load_preset(preset_id or document.front.preset)
    sources = load_sources(document, input_path.parent, input_path.name)
    prepared = prepare(document, preset, sources, name=input_path.name)
    target = output or input_path.with_name(f"{input_path.stem}_gost.docx")
    if target.resolve() == input_path.resolve():
        raise DocumentError("Результат не может перезаписать исходный файл", str(input_path))
    data = render(prepared, preset, base_dir=input_path.parent, name=input_path.name, document=document)
    target.write_bytes(data)
    return BuildResult(target, prepared.stats, prepared.warnings)


def load_sources(document: ParsedDocument, base_dir: Path, name: str) -> list[Source]:
    """Sources from the front matter: a list, or a path to a JSON/YAML file (relative to the document)."""
    raw: Any = document.front.sources
    if isinstance(raw, str):
        path = base_dir / raw
        data = read_data(str(path))
        raw = data.get("sources") if isinstance(data, dict) else data
        if not isinstance(raw, list):
            raise DocumentError(f"В файле «{path.name}» ожидался список источников", str(path))
    items = resolve(raw, location=name)
    if items.unresolved:
        details = "; ".join(f"{item.index}. «{item.input}»: {item.reason}" for item in items.unresolved)
        raise DocumentError(f"Не удалось получить данные источников: {details}", name)
    sources = items.sources
    for source in sources:
        if not source.id:
            raise DocumentError(f"У источника «{source.title}» нет id — на него нельзя сослаться", name)
    return sources


def render(
    prepared: Prepared, preset: Preset, *, base_dir: Path, name: str, document: ParsedDocument | None = None
) -> bytes:
    doc = docx.Document()
    styles.apply_page(doc, preset)
    styles.apply_styles(doc, preset)
    _page_numbers(doc, preset)
    writer = _Writer(doc, preset, base_dir, name)
    for block in prepared.blocks:
        writer.block(block)
    _update_fields_on_open(doc)
    title = (document.front.title_page.get("title") if document else None) or ""
    _fixed_properties(doc, str(title))
    buffer = io.BytesIO()
    doc.save(buffer)
    return _normalize_zip(buffer.getvalue())


class _Writer:
    def __init__(self, doc: Document, preset: Preset, base_dir: Path, name: str) -> None:
        self.doc = doc
        self.preset = preset
        self.base_dir = base_dir
        self.name = name

    def block(self, block: OutBlock) -> None:
        if isinstance(block, OutHeading):
            self.heading(block)
        elif isinstance(block, OutParagraph):
            _inline(self.doc.add_paragraph(), block.text)
        elif isinstance(block, OutList):
            for item in block.items:
                _inline(self.doc.add_paragraph(), item)
        elif isinstance(block, OutFigure):
            self.figure(block)
        elif isinstance(block, OutTable):
            self.table(block)
        elif isinstance(block, OutFormula):
            self.formula(block)
        elif isinstance(block, OutCode):
            for line in block.lines or [""]:
                self.doc.add_paragraph(line, style=styles.CODE)
        else:
            assert isinstance(block, OutBibliography)
            for entry in block.entries:
                self.doc.add_paragraph(entry)

    def heading(self, block: OutHeading) -> None:
        headings = self.preset.headings
        if block.kind == "toc":
            paragraph = self.doc.add_paragraph()
            styles.heading_format(paragraph, headings.structural, self.preset)
            run = paragraph.add_run(block.text)
            run.bold = headings.structural.bold
            _toc_field(self.doc.add_paragraph())
            return
        if block.kind in ("structural", "bibliography"):
            paragraph = self.doc.add_paragraph(block.text, style="Heading 1")
            styles.heading_format(paragraph, headings.structural, self.preset)
            return
        if block.kind == "appendix":
            paragraph = self.doc.add_paragraph(style="Heading 1")
            styles.heading_format(paragraph, headings.structural, self.preset)
            paragraph.paragraph_format.page_break_before = self.preset.appendices.new_page
            paragraph.add_run(block.text)
            if block.appendix_title:
                # Заголовок приложения — отдельной строкой по центру (ГОСТ 7.32-2017, п. 6.17.3).
                paragraph.add_run().add_break(WD_BREAK.LINE)
                paragraph.add_run(block.appendix_title)
            return
        style = headings.level1 if block.level == 1 else headings.level2
        paragraph = self.doc.add_paragraph(block.text, style=f"Heading {min(block.level, 3)}")
        styles.heading_format(paragraph, style, self.preset)

    def figure(self, block: OutFigure) -> None:
        path = self.base_dir / block.path
        if not path.is_file():
            raise DocumentError(f"Не найден файл рисунка «{block.path}»", f"{self.name}:{block.line}")
        try:
            image = DocxImage.from_file(str(path))
        except UnrecognizedImageError:
            raise DocumentError(
                f"Формат рисунка «{block.path}» не поддерживается: используйте PNG, JPG, GIF, BMP или TIFF",
                f"{self.name}:{block.line}",
            ) from None
        max_width = Mm(styles.text_width_mm(self.preset))
        width = Emu(min(image.width, max_width))
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles.no_indent(paragraph)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=width)
        caption = self.doc.add_paragraph(style="Caption")
        caption.alignment = styles.ALIGN[self.preset.figures.caption_align]
        caption.paragraph_format.line_spacing = self.preset.figures.caption_line_spacing
        _inline(caption, block.caption)

    def table(self, block: OutTable) -> None:
        caption = self.doc.add_paragraph(style="Caption")
        caption.alignment = styles.ALIGN[self.preset.tables.caption_align]
        caption.paragraph_format.line_spacing = self.preset.tables.caption_line_spacing
        caption.paragraph_format.keep_with_next = True
        _inline(caption, block.caption)
        table = self.doc.add_table(rows=1 + len(block.rows), cols=len(block.header))
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for column, text in enumerate(block.header):
            cell = table.cell(0, column).paragraphs[0]
            cell.style = self.doc.styles[styles.TABLE_TEXT]
            cell.alignment = WD_ALIGN_PARAGRAPH.CENTER  # заголовки граф по центру (п. 6.6.6)
            _inline(cell, text)
        for row_index, row in enumerate(block.rows, start=1):
            for column, text in enumerate(row):
                cell = table.cell(row_index, column).paragraphs[0]
                cell.style = self.doc.styles[styles.TABLE_TEXT]
                _inline(cell, text)
        _repeat_header(table)

    def formula(self, block: OutFormula) -> None:
        try:
            math = latex_to_omml(block.latex)
        except MathError as exc:
            raise DocumentError(f"Формула не распознана: {exc}", f"{self.name}:{block.line}") from None
        paragraph = self.doc.add_paragraph()
        styles.no_indent(paragraph)
        fmt = paragraph.paragraph_format
        if self.preset.formulas.blank_line_around:
            # Свободная строка выше и ниже формулы (ГОСТ 7.32-2017, п. 6.8.1) — интервалом, не пустым абзацем.
            line = Pt(self.preset.text.size_pt * self.preset.text.line_spacing)
            fmt.space_before = line
            fmt.space_after = line
        if block.number is None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph._p.append(math)
            return
        width = Mm(styles.text_width_mm(self.preset))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.tab_stops.add_tab_stop(Emu(width // 2), WD_TAB_ALIGNMENT.CENTER)
        fmt.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run("\t")
        paragraph._p.append(math)
        paragraph.add_run(f"\t{block.number}")


def _inline(paragraph: Paragraph, text: str) -> None:
    """**полужирный**, *курсив*, `моноширинный`; перевод строки — разрыв строки; остальное — как есть."""
    for number, line in enumerate(text.split("\n")):
        if number:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        _inline_line(paragraph, line)


def _inline_line(paragraph: Paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = styles.CODE_FONT
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _field_run(paragraph: Paragraph, kind: str) -> None:
    run = paragraph.add_run()
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    run._r.append(element)


def _instruction(paragraph: Paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    element = OxmlElement("w:instrText")
    element.set(qn("xml:space"), "preserve")
    element.text = f" {instruction} "
    run._r.append(element)


def _toc_field(paragraph: Paragraph) -> None:
    styles.no_indent(paragraph)
    _field_run(paragraph, "begin")
    _instruction(paragraph, TOC_FIELD)
    _field_run(paragraph, "separate")
    paragraph.add_run(TOC_PLACEHOLDER)
    _field_run(paragraph, "end")


def _page_numbers(doc: Document, preset: Preset) -> None:
    position = preset.page_numbers.position
    for section in doc.sections:
        part = section.footer if position.startswith("bottom") else section.header
        part.is_linked_to_previous = False
        paragraph = part.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if position.endswith("center") else WD_ALIGN_PARAGRAPH.RIGHT
        )
        styles.no_indent(paragraph)
        paragraph.paragraph_format.line_spacing = 1.0
        _field_run(paragraph, "begin")
        _instruction(paragraph, "PAGE")
        _field_run(paragraph, "separate")
        paragraph.add_run("1")
        _field_run(paragraph, "end")


def _repeat_header(table: Table) -> None:
    """Шапка таблицы повторяется на каждой странице при переносе."""
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _update_fields_on_open(doc: Document) -> None:
    update = ooxml.element("w:updateFields", val="true")
    ooxml.insert_before_successors(doc.settings.element, update, ooxml.AFTER_UPDATE_FIELDS)


def _fixed_properties(doc: Document, title: str) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = title
    props.revision = 1
    props.created = FIXED_TIME.replace(tzinfo=None)
    props.modified = FIXED_TIME.replace(tzinfo=None)
    props.last_printed = FIXED_TIME.replace(tzinfo=None)


def _normalize_zip(data: bytes) -> bytes:
    """Rewrite the package with fixed timestamps so the same input gives the same bytes."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(item.filename))
    return out.getvalue()
